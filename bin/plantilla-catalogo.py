"""
Catalog import kit for Frave (phase 6):

- PLANTILLA_IMPORTACION_WOOCOMMERCE.xlsx: the sheet to fill in ("Importar a Woo", with the
  native WooCommerce headers so the importer maps them by itself), filled examples, the valid
  category paths, the attributes with their suggested values and the instructions.
- bin/muestra-catalogo.csv: ten [MUESTRA] products (one variable, with variations) to try the
  import and the catalog locally. Never import it in production.

    pip install openpyxl
    python bin/plantilla-catalogo.py

The category tree mirrors bin/crear-lineas.php and the attributes bin/crear-atributos.php.
"""
import csv
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.join(os.path.dirname(__file__), '..')
IMAGE_BASE = 'https://frave.pe/wp-content/themes/frave/assets/media/lineas/'

TREE = {
    'Sabores': {'Bebidas y productos lácteos': [], 'Confitería y panadería': [], 'Salados': [], 'Cuidado bucal': []},
    'Fragancias': {
        'Perfumería': ['Perfumería de diseñador', 'Perfumería árabe', 'Perfumería exclusiva (nicho)'],
        'Fragancias cosméticas': ['Cuidado del cabello', 'Cuidado de la piel', 'Cuidado de mascotas'],
        'Fragancias para el cuidado del hogar': ['Cuidado textil', 'Cuidado de superficies', 'Cocina y lavavajillas'],
        'Ambientación': ['Difusores de varitas', 'Ambientación de vehículos', 'Aerosoles y sprays', 'Marketing olfativo'],
    },
    'Envases': {'Frascos de vidrio para perfumería': [], 'Frascos para difusores': [], 'Válvulas': []},
}

ATTRIBUTES = [
    ('Presentación', 'Todas', 'Formatos de venta: 250 g, 500 g, 1 kg, 5 kg… Escribe siempre igual (número, espacio, unidad).'),
    ('Familia olfativa', 'Fragancias', 'Cítrica, Floral, Frutal, Amaderada, Oriental, Aromática, Gourmand, Fresca, Especiada'),
    ('Perfil de sabor', 'Sabores', 'Frutal, Cítrico, Lácteo, Dulce, Tostado, Especiado, Salado, Mentolado'),
    ('Aplicación', 'Fragancias y sabores', 'Perfumería fina, Body splash, Cremas y lociones, Jabones, Shampoo y acondicionador, Detergentes y suavizantes, Limpiadores, Velas, Difusores, Ambientadores, Bebidas, Lácteos, Confitería, Panadería, Snacks, Cuidado bucal'),
    ('Solubilidad', 'Fragancias y sabores', 'Hidrosoluble, Liposoluble'),
    ('Capacidad', 'Envases', '30 ml, 50 ml, 100 ml, 120 ml…'),
    ('Material', 'Envases', 'Vidrio, PET, Polipropileno, Aluminio'),
    ('Cierre', 'Envases', 'Válvula atomizadora, Bomba dosificadora, Tapa rosca, Tapa a presión, Gotero'),
]

HEADERS = ['Type', 'SKU', 'Name', 'Published', 'Is featured?', 'Visibility in catalog', 'Short description', 'Description',
           'In stock?', 'Regular price', 'Categories', 'Images', 'Parent', 'Position']
for n in range(1, 5):
    HEADERS += [f'Attribute {n} name', f'Attribute {n} value(s)', f'Attribute {n} visible', f'Attribute {n} global', f'Attribute {n} default']

GUIDE = [
    ('Type', 'Sí', 'simple (un producto), variable (producto con opciones, p. ej. capacidades) o variation (cada opción de un variable, debajo de su padre).', 'simple'),
    ('SKU', 'Sí', 'Código único y estable. Las variaciones usan el del padre más un sufijo.', 'FRA-ARA-001'),
    ('Name', 'Sí', 'Nombre comercial. Deja vacío en las variaciones: WooCommerce lo arma solo.', 'Fragancia Oud Ámbar'),
    ('Published', 'Sí', '1 publicado, 0 privado, -1 borrador. Prueba primero con -1.', '1'),
    ('Is featured?', 'No', '1 para que aparezca en «Productos destacados» de la portada y de su línea.', '0'),
    ('Visibility in catalog', 'No', 'visible (por defecto), catalog, search o hidden.', 'visible'),
    ('Short description', 'Sí', 'Dos o tres líneas: qué es y para qué sirve. Se ve junto al título y en la vista previa al compartir.', 'Fragancia intensa de oud y ámbar…'),
    ('Description', 'No', 'Texto completo: notas, recomendaciones de uso, dosis sugerida, presentación.', ''),
    ('In stock?', 'No', '1 disponible, 0 agotado. En modo catálogo solo informa.', '1'),
    ('Regular price', 'No', 'Precio en soles, sin símbolo. En modo catálogo no se muestra; sirve para la futura tienda.', ''),
    ('Categories', 'Sí', 'Ruta completa de la hoja «Categorías», con « > » entre niveles. Varias, separadas por comas.', 'Fragancias > Perfumería > Perfumería árabe'),
    ('Images', 'Sí', 'URL de la foto (o varias separadas por comas). La primera es la principal. Sube antes las fotos a Medios y copia su URL.', 'https://frave.pe/wp-content/uploads/…/oud-ambar.jpg'),
    ('Parent', 'Solo variaciones', 'SKU del producto variable al que pertenece la variación.', 'ENV-VID-001'),
    ('Position', 'No', 'Orden de la variación dentro del producto (0, 1, 2…).', '0'),
    ('Attribute N name', 'No', 'Nombre exacto de la hoja «Atributos». Hasta 4 por producto.', 'Familia olfativa'),
    ('Attribute N value(s)', 'No', 'Valores separados por comas (en una variación, uno solo).', 'Oriental, Amaderada'),
    ('Attribute N visible', 'No', '1 para mostrarlo en la ficha («Información adicional»).', '1'),
    ('Attribute N global', 'No', 'Siempre 1: los atributos globales son los que aparecen en los filtros.', '1'),
    ('Attribute N default', 'No', 'Solo en productos variables: la opción elegida al abrir la ficha.', '50 ml'),
]


def attrs(*pairs):
    """Up to four (name, values, default) tuples as the 20 attribute cells."""
    cells = []
    for n in range(4):
        if n < len(pairs):
            name, values, default = (pairs[n] + ('',))[:3]
            cells += [name, values, '1', '1', default]
        else:
            cells += ['', '', '', '', '']
    return cells


def row(kind, sku, name, short, description, categories, image, price='', featured='0', parent='', position='', attributes=()):
    images = IMAGE_BASE + image if image and not image.startswith('http') else image
    return [kind, sku, name, '1', featured, 'visible', short, description, '1', price, categories, images, parent, position] + attrs(*attributes)


SAMPLES = [
    row('simple', 'FRA-ARA-001', '[MUESTRA] Fragancia Oud Ámbar', 'Fragancia intensa y duradera de oud, ámbar y almizcle, inspirada en la perfumería árabe.',
        'Composición oriental amaderada con salida especiada, corazón de rosa y fondo de oud, ámbar y almizcle. Recomendada para perfumería fina y body splash.',
        'Fragancias > Perfumería > Perfumería árabe', 'fragancias-tarjeta.webp', '180', '1',
        attributes=(('Familia olfativa', 'Oriental, Amaderada'), ('Aplicación', 'Perfumería fina, Body splash'), ('Solubilidad', 'Liposoluble'), ('Presentación', '250 g, 500 g, 1 kg'))),
    row('simple', 'FRA-COS-001', '[MUESTRA] Fragancia Coco Vainilla', 'Fragancia gourmand cremosa para cremas, lociones y jabones.',
        'Notas de coco, vainilla y almendra, compatible con emulsiones y bases jabonosas.',
        'Fragancias > Fragancias cosméticas > Cuidado de la piel', 'fragancias-tarjeta.webp', '120',
        attributes=(('Familia olfativa', 'Gourmand'), ('Aplicación', 'Cremas y lociones, Jabones, Body splash'), ('Solubilidad', 'Hidrosoluble'), ('Presentación', '500 g, 1 kg'))),
    row('simple', 'FRA-HOG-001', '[MUESTRA] Fragancia Lavanda Fresca', 'Fragancia aromática y persistente para detergentes y suavizantes.',
        'Lavanda, notas verdes y almizcle limpio, pensada para cuidado textil y limpiadores multiusos.',
        'Fragancias > Fragancias para el cuidado del hogar > Cuidado textil', 'fragancias-tarjeta.webp', '95',
        attributes=(('Familia olfativa', 'Aromática, Fresca'), ('Aplicación', 'Detergentes y suavizantes, Limpiadores'), ('Solubilidad', 'Hidrosoluble'), ('Presentación', '1 kg, 5 kg'))),
    row('simple', 'FRA-AMB-001', '[MUESTRA] Fragancia Bambú y Té Verde', 'Aroma fresco y envolvente para difusores, velas y ambientadores.',
        'Bambú, té verde y cítricos suaves para ambientes relajantes y espacios comerciales.',
        'Fragancias > Ambientación > Difusores de varitas', 'fragancias-tarjeta.webp', '110', '1',
        attributes=(('Familia olfativa', 'Fresca, Aromática'), ('Aplicación', 'Difusores, Velas, Ambientadores'), ('Solubilidad', 'Liposoluble'), ('Presentación', '500 g, 1 kg'))),
    row('simple', 'SAB-BEB-001', '[MUESTRA] Sabor Maracuyá', 'Sabor frutal y tropical para bebidas, néctares y yogures.',
        'Perfil de maracuyá madura con acidez natural, estable en bebidas y productos lácteos.',
        'Sabores > Bebidas y productos lácteos', 'sabores-tarjeta.webp', '85', '1',
        attributes=(('Perfil de sabor', 'Frutal, Cítrico'), ('Aplicación', 'Bebidas, Lácteos'), ('Solubilidad', 'Hidrosoluble'), ('Presentación', '1 kg, 5 kg'))),
    row('simple', 'SAB-PAN-001', '[MUESTRA] Sabor Vainilla Panetón', 'Sabor dulce y lácteo para panetones, kekes y confitería.',
        'Vainilla con notas de mantequilla y frutas confitadas, resistente al horneado.',
        'Sabores > Confitería y panadería', 'sabores-tarjeta.webp', '75',
        attributes=(('Perfil de sabor', 'Dulce, Lácteo'), ('Aplicación', 'Panadería, Confitería'), ('Solubilidad', 'Liposoluble'), ('Presentación', '1 kg, 5 kg'))),
    row('simple', 'SAB-BUC-001', '[MUESTRA] Sabor Menta Fresca', 'Sabor mentolado y refrescante para pastas dentales y enjuagues.',
        'Menta piperita con un toque dulce, para productos de higiene bucal.',
        'Sabores > Cuidado bucal', 'sabores-tarjeta.webp', '90',
        attributes=(('Perfil de sabor', 'Mentolado'), ('Aplicación', 'Cuidado bucal'), ('Solubilidad', 'Hidrosoluble'), ('Presentación', '1 kg'))),
    row('variable', 'ENV-VID-001', '[MUESTRA] Frasco de vidrio cuadrado', 'Frasco de vidrio transparente con válvula atomizadora, en tres capacidades.',
        'Vidrio grueso de base cuadrada, para perfumes y colonias. Incluye válvula atomizadora plateada.',
        'Envases > Frascos de vidrio para perfumería', 'envases-tarjeta.webp', '', '1',
        attributes=(('Capacidad', '30 ml, 50 ml, 100 ml', '50 ml'), ('Material', 'Vidrio'), ('Cierre', 'Válvula atomizadora'))),
]
for position, (capacity, price) in enumerate((('30 ml', '6.5'), ('50 ml', '8'), ('100 ml', '11'))):
    variation = row('variation', f'ENV-VID-001-{capacity.split()[0]}', '', '', '', '', '', price, parent='ENV-VID-001', position=str(position),
                    attributes=(('Capacidad', capacity),))
    variation[3] = '1'
    SAMPLES.append(variation)
SAMPLES += [
    row('simple', 'ENV-DIF-001', '[MUESTRA] Frasco para difusor 120 ml', 'Frasco de vidrio para difusores de varitas, con tapa rosca.',
        'Boca ancha para varitas de ratán, ideal para ambientación de hogares y espacios comerciales.',
        'Envases > Frascos para difusores', 'envases-tarjeta.webp', '9',
        attributes=(('Capacidad', '120 ml'), ('Material', 'Vidrio'), ('Cierre', 'Tapa rosca'))),
    row('simple', 'ENV-VAL-001', '[MUESTRA] Válvula atomizadora 15 mm', 'Válvula atomizadora de engaste para frascos de perfumería.',
        'Rosca de 15 mm con atomización fina y uniforme.',
        'Envases > Válvulas', 'envases-tarjeta.webp', '2.5',
        attributes=(('Material', 'Aluminio'), ('Cierre', 'Válvula atomizadora'))),
]


def category_paths():
    paths = []
    for line, groups in TREE.items():
        paths.append(line)
        for group, children in groups.items():
            paths.append(f'{line} > {group}')
            paths += [f'{line} > {group} > {child}' for child in children]
    return paths


HEAD_FILL = PatternFill('solid', fgColor='FC4C02')
HEAD_FONT = Font(bold=True, color='FFFFFF')


def style_header(ws, widths):
    for cell in ws[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    ws.freeze_panes = 'A2'
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, index).column_letter].width = width


wb = Workbook()
importer = wb.active
importer.title = 'Importar a Woo'
importer.append(HEADERS)
style_header(importer, [10, 16, 34, 10, 10, 14, 40, 40, 9, 11, 42, 40, 14, 9] + [18, 26, 9, 9, 12] * 4)

examples = wb.create_sheet('Ejemplos')
examples.append(HEADERS)
for sample in SAMPLES:
    examples.append(sample)
style_header(examples, [10, 16, 34, 10, 10, 14, 40, 40, 9, 11, 42, 40, 14, 9] + [18, 26, 9, 9, 12] * 4)

categories = wb.create_sheet('Categorías')
categories.append(['Ruta para la columna Categories'])
for path in category_paths():
    categories.append([path])
style_header(categories, [70])

attributes = wb.create_sheet('Atributos')
attributes.append(['Nombre (Attribute N name)', 'Línea', 'Valores sugeridos (Attribute N value(s))'])
for item in ATTRIBUTES:
    attributes.append(list(item))
style_header(attributes, [26, 22, 110])

guide = wb.create_sheet('Instrucciones')
guide.append(['Columna', 'Obligatoria', 'Qué escribir', 'Ejemplo'])
for item in GUIDE:
    guide.append(list(item))
guide.append([])
for line in (
    'Completa solo la hoja «Importar a Woo» y expórtala como CSV UTF-8 (Archivo → Guardar como → CSV UTF-8).',
    'Importa el CSV en WooCommerce → Productos → Importar. Las columnas se asignan solas; revisa la pantalla de asignación y continúa.',
    'Antes de importar: crea las categorías (bin/crear-lineas.php) y los atributos (bin/crear-atributos.php), y sube las fotos a Medios.',
    'Prueba primero con dos o tres productos en borrador (Published = -1). Para corregir, vuelve a importar con «Actualizar productos existentes» (se identifican por SKU).',
    'Fotos: cuadradas o 4:5, mínimo 1200 px, fondo marfil liso como el resto del sitio, JPG o WebP de menos de 300 KB. Ver GUIA_CATALOGO.md.',
):
    guide.append([line])
style_header(guide, [24, 14, 90, 40])
for row_cells in guide.iter_rows(min_row=2):
    for cell in row_cells:
        cell.alignment = Alignment(vertical='top', wrap_text=True)

# Drop-downs on the import sheet (warnings, so other valid values are still allowed).
last = len(category_paths()) + 1
validations = [
    ('A', '"simple,variable,variation"'),
    ('D', '"1,0,-1"'),
    ('E', '"0,1"'),
    ('F', '"visible,catalog,search,hidden"'),
    ('I', '"1,0"'),
    ('K', f"'Categorías'!$A$2:$A${last}"),
]
for column, formula in validations:
    rule = DataValidation(type='list', formula1=formula, allow_blank=True, showErrorMessage=True, errorStyle='warning')
    importer.add_data_validation(rule)
    rule.add(f'{column}2:{column}1000')
names = DataValidation(type='list', formula1=f"'Atributos'!$A$2:$A${len(ATTRIBUTES) + 1}", allow_blank=True, showErrorMessage=True, errorStyle='warning')
importer.add_data_validation(names)
for n in range(4):
    column = importer.cell(1, 15 + n * 5).column_letter
    names.add(f'{column}2:{column}1000')

wb.save(os.path.join(ROOT, 'PLANTILLA_IMPORTACION_WOOCOMMERCE.xlsx'))
with open(os.path.join(os.path.dirname(__file__), 'muestra-catalogo.csv'), 'w', encoding='utf-8-sig', newline='') as handle:
    writer = csv.writer(handle)
    writer.writerow(HEADERS)
    writer.writerows(SAMPLES)
print(f'{len(SAMPLES)} filas de muestra, {len(category_paths())} rutas de categoría')
